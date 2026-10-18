import pptx
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])

# Add text box
txBox = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(8), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "Hello World"
p.font.size = Pt(36)
p.font.color.rgb = RGBColor(200, 30, 30)

# Add transition
transition_xml = (
    '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'spd="med" advClick="1">'
    '<p:fade/>'
    '</p:transition>'
)
slide._element.append(parse_xml(transition_xml))

# Add build animation on the shape
shape_id = txBox.shape_id
timing_xml = f'''
<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="always" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                <p:par>
                  <p:cTn id="3" fill="hold">
                    <p:stCondLst>
                      <p:cond delay="0"/>
                    </p:stCondLst>
                    <p:childTnLst>
                      <p:par>
                        <p:cTn id="4" fill="hold">
                          <p:stCondLst>
                            <p:cond delay="0"/>
                          </p:stCondLst>
                          <p:childTnLst>
                            <p:set>
                              <p:cBhvr>
                                <p:cTn id="5" dur="1" fill="hold"/>
                                <p:tgtEl>
                                  <p:spTgt spid="{shape_id}"/>
                                </p:tgtEl>
                                <p:attrNameLst>
                                  <p:attrName>style.visibility</p:attrName>
                                </p:attrNameLst>
                              </p:cBhvr>
                              <p:to>
                                <p:strVal val="visible"/>
                              </p:to>
                            </p:set>
                            <p:animEffect transition="in" filter="fade">
                              <p:cBhvr>
                                <p:cTn id="6" dur="500"/>
                                <p:tgtEl>
                                  <p:spTgt spid="{shape_id}"/>
                                </p:tgtEl>
                              </p:cBhvr>
                            </p:animEffect>
                          </p:childTnLst>
                        </p:cTn>
                      </p:par>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>
'''
slide._element.append(parse_xml(timing_xml))

prs.save("test_out.pptx")
print("Saved test_out.pptx successfully!")
